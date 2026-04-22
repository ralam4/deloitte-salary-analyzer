"""Recompute src/data/salaryData.js from the 2025 Deloitte Salary Survey XLSX.

Base salary benchmarks are computed against the FY26 Base Salary column
(respondents' post-June-2025 salaries). Raise and promotion stats are
computed against FY25->FY26 percent change, unchanged from the prior run.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import numpy as np
import openpyxl

REPO_ROOT = Path(__file__).resolve().parents[1]
XLSX_PATH = REPO_ROOT / "2025 Deloitte Salary Survey Responses.xlsx"
OUT_PATH = REPO_ROOT / "src" / "data" / "salaryData.js"


HEADER = '''// Stats recomputed from 2025 Deloitte Salary Survey Responses.xlsx
// Base salary benchmarks are computed against the FY26 Base Salary column
// (respondents' post-June-2025 salaries).
// Raise stats are FY25->FY26 percent change.
// Regenerate with: python3 scripts/recompute_stats_fy26.py
'''


def as_js(name: str, value) -> str:
    return f"export const {name} = {json.dumps(value, indent=2)};\n\n"


def write_js(out: dict) -> None:
    parts = [HEADER + "\n"]
    parts.append(as_js("LEVEL_STATS", out["level"]))
    parts.append(as_js("GPS_COMMERCIAL_STATS", out["gps"]))
    parts.append(as_js("BUSINESS_STATS", out["biz"]))
    parts.append(as_js("PORTFOLIO_STATS", out["port"]))
    parts.append(as_js("USDC_STATS", out["usdc"]))
    parts.append(as_js("MBA_STATS", out["mba"]))
    parts.append(as_js("MBA_PREMIUM", out["premium"]))
    parts.append(as_js("YEARS_AT_LEVEL_MANAGER", out["years"]))
    parts.append(as_js("PROMOTION_RAISES", out["promos"]))
    parts.append(as_js("NON_PROMOTION_RAISE", out["non_promo"]))

    parts.append("export const LEVELS = Object.keys(LEVEL_STATS);\n")
    parts.append('export const BUSINESSES = ["Consulting Services", "Audit & Assurance", "Tax", "Enabling Areas"];\n')
    parts.append("export const PORTFOLIOS = " + json.dumps(PORTFOLIO_ORDER) + ";\n")
    parts.append('export const GPS_COMM = ["Commercial", "GPS"];\n')
    parts.append('export const EDUCATION_LEVELS = ["Bachelor\\u0027s", "Non-MBA Master\\u0027s", "MBA", "PhD / Other"];\n')
    parts.append('export const BUSINESS_MODELS = ["Core (Traditional)", "USDC"];\n')
    parts.append("export const NEXT_LEVEL = " + json.dumps({
        "Analyst / Jr Staff": "Consultant / Staff",
        "Consultant / Staff": "Senior Consultant / Specialist Senior / Senior",
        "Senior Consultant / Specialist Senior / Senior": "Manager / Specialist Master",
        "Manager / Specialist Master": "Senior Manager / Specialist Leader",
    }) + ";\n\n")
    parts.append(f"export const totalRespondents = {out['count']};\n")

    OUT_PATH.write_text("".join(parts))


def load_rows() -> list[dict]:
    wb = openpyxl.load_workbook(XLSX_PATH, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for raw in ws.iter_rows(min_row=2, values_only=True):
        rows.append(dict(zip(headers, raw)))
    return rows


LEVEL_ORDER = [
    "Analyst / Jr Staff",
    "Consultant / Staff",
    "Senior Consultant / Specialist Senior / Senior",
    "Manager / Specialist Master",
    "Senior Manager / Specialist Leader",
]


def normalize_business(raw: str | None) -> str | None:
    if not raw:
        return None
    s = raw.strip()
    if s in ("Consulting", "Advisory"):
        return "Consulting Services"
    if s in ("Consulting Services", "Audit & Assurance", "Tax", "Enabling Areas"):
        return s
    return None


def clean(rows: list[dict]) -> list[dict]:
    out = []
    for r in rows:
        fy26 = r.get("FY26 Base Salary (USD)")
        level = r.get("FY26 Level")
        if not isinstance(fy26, (int, float)) or fy26 <= 0:
            continue
        if level not in LEVEL_ORDER:
            continue
        if r.get("Data Quality Concern"):
            continue
        out.append(r)
    return out


def pct_block(values: list[float], *, include_p10_p90: bool = True) -> dict:
    arr = np.array([v for v in values if isinstance(v, (int, float)) and v > 0], dtype=float)
    if len(arr) == 0:
        return {}
    out = {
        "p25": float(np.percentile(arr, 25)),
        "p50": float(np.percentile(arr, 50)),
        "p75": float(np.percentile(arr, 75)),
        "mean": float(arr.mean()),
    }
    if include_p10_p90:
        out["p10"] = float(np.percentile(arr, 10))
        out["p90"] = float(np.percentile(arr, 90))
    return out


def salary_stats(rows: list[dict]) -> dict:
    salary = pct_block([r["FY26 Base Salary (USD)"] for r in rows])
    aip = pct_block([r.get("AIP (USD)") or 0 for r in rows], include_p10_p90=False)
    tc_values = [
        (r["FY26 Base Salary (USD)"] or 0) + (r.get("AIP (USD)") or 0)
        for r in rows
    ]
    tc = pct_block(tc_values, include_p10_p90=False)
    return {
        "count": len(rows),
        "salary": {k: round(v) for k, v in salary.items()},
        "aip": {k: round(v) for k, v in aip.items()},
        "tc": {k: round(v) for k, v in tc.items() if k in ("p25", "p50", "p75")},
    }


def level_stats(rows: list[dict]) -> dict:
    out = {}
    for lvl in LEVEL_ORDER:
        subset = [r for r in rows if r["FY26 Level"] == lvl]
        out[lvl] = salary_stats(subset)
        assert out[lvl]["salary"]["p25"] <= out[lvl]["salary"]["p50"] <= out[lvl]["salary"]["p75"], lvl
    return out


PORTFOLIO_ORDER = [
    "AI & Engineering",
    "Strategy & Transactions",
    "Customer",
    "Human Capital",
    "Enterprise Performance",
    "Cyber",
    "Finance Transformation",
    "Regulatory, Risk & Forensic",
    "Other",
]
MIN_N_PORTFOLIO = 30


def business_stats(rows: list[dict]) -> dict:
    out = {}
    for lvl in LEVEL_ORDER:
        lvl_rows = [r for r in rows if r["FY26 Level"] == lvl]
        out[lvl] = {}
        for biz in ("Consulting Services", "Audit & Assurance", "Tax", "Enabling Areas"):
            subset = [r for r in lvl_rows if normalize_business(r.get("FY26 Global Business")) == biz]
            if len(subset) >= 10:
                out[lvl][biz] = salary_stats(subset)
    return out


def gps_commercial_stats(rows: list[dict]) -> dict:
    out = {}
    for lvl in LEVEL_ORDER:
        lvl_rows = [r for r in rows if r["FY26 Level"] == lvl]
        out[lvl] = {}
        for key in ("GPS", "Commercial"):
            subset = [r for r in lvl_rows if (r.get("GPS or Commercial") or "").strip() == key]
            if len(subset) >= 10:
                out[lvl][key] = salary_stats(subset)
    return out


def portfolio_stats(rows: list[dict]) -> dict:
    out = {}
    for lvl in LEVEL_ORDER:
        lvl_rows = [r for r in rows if r["FY26 Level"] == lvl]
        level_block = {}
        for p in PORTFOLIO_ORDER:
            subset = [r for r in lvl_rows if (r.get("FY26 Offering Portfolio") or "").strip() == p]
            if len(subset) >= MIN_N_PORTFOLIO:
                level_block[p] = salary_stats(subset)
        if level_block:
            out[lvl] = level_block
    return out


def usdc_stats(rows: list[dict]) -> dict:
    out = {}
    for lvl in LEVEL_ORDER[:4]:  # no USDC data at Senior Manager in prior output
        lvl_rows = [r for r in rows if r["FY26 Level"] == lvl]
        block = {}
        for key, match in (("USDC", "USDC"), ("Core", "Core (Traditional)")):
            subset = [r for r in lvl_rows if (r.get("Talent Model") or "").strip() == match]
            if len(subset) >= 10:
                salary = pct_block(
                    [r["FY26 Base Salary (USD)"] for r in subset],
                    include_p10_p90=False,
                )
                block[key] = {
                    "count": len(subset),
                    "salary": {k: round(v) for k, v in salary.items() if k in ("p25", "p50", "p75")},
                }
        if "USDC" in block and "Core" in block:
            out[lvl] = block
    return out


def mba_stats(rows: list[dict]) -> dict:
    out = {}
    for lvl in LEVEL_ORDER[1:]:  # MBA stats skip Analyst
        lvl_rows = [r for r in rows if r["FY26 Level"] == lvl]
        mba = [r for r in lvl_rows if (r.get("Education Level") or "").strip() == "MBA"]
        non = [r for r in lvl_rows if (r.get("Education Level") or "").strip() and (r.get("Education Level") or "").strip() != "MBA"]
        if len(mba) >= 20 and len(non) >= 20:
            out[lvl] = {"MBA": salary_stats(mba), "NonMBA": salary_stats(non)}
    return out


def mba_premium(rows: list[dict]) -> dict:
    out = {}
    for lvl in LEVEL_ORDER[1:]:
        lvl_rows = [r for r in rows if r["FY26 Level"] == lvl]
        mba = [r["FY26 Base Salary (USD)"] for r in lvl_rows if (r.get("Education Level") or "").strip() == "MBA"]
        bach = [r["FY26 Base Salary (USD)"] for r in lvl_rows if (r.get("Education Level") or "").strip() == "Bachelor's"]
        if len(mba) >= 20 and len(bach) >= 20:
            m_med = float(np.median(mba))
            b_med = float(np.median(bach))
            out[lvl] = {
                "mba": {"n": len(mba), "median": round(m_med)},
                "bachelors": {"n": len(bach), "median": round(b_med)},
                "delta": round(m_med - b_med),
                "deltaPct": round((m_med - b_med) / b_med, 3),
            }
    return out


def years_at_level_manager(rows: list[dict]) -> dict:
    mgr = [r for r in rows if r["FY26 Level"] == "Manager / Specialist Master"]
    out = {}
    for y in (1, 2, 3, 4, 5):
        subset = [r for r in mgr if r.get("Rounded Years at level") == y]
        if len(subset) >= 10:
            out[y] = {
                "n": len(subset),
                "median": round(float(np.median([r["FY26 Base Salary (USD)"] for r in subset]))),
            }
    return out


LEVEL_TO_PREV = {
    "Consultant / Staff": "Analyst / Jr Staff",
    "Senior Consultant / Specialist Senior / Senior": "Consultant / Staff",
    "Manager / Specialist Master": "Senior Consultant / Specialist Senior / Senior",
    "Senior Manager / Specialist Leader": "Manager / Specialist Master",
}
FROM_LABEL = {
    "Consultant / Staff": ("Analyst", "Consultant"),
    "Senior Consultant / Specialist Senior / Senior": ("Consultant", "Senior Consultant"),
    "Manager / Specialist Master": ("Senior Consultant", "Manager"),
    "Senior Manager / Specialist Leader": ("Manager", "Senior Manager"),
}


def raise_pct(r: dict) -> float | None:
    fy25 = r.get("FY25 Base Salary (USD)")
    fy26 = r.get("FY26 Base Salary (USD)")
    if not (isinstance(fy25, (int, float)) and fy25 > 0 and isinstance(fy26, (int, float)) and fy26 > 0):
        return None
    return (fy26 - fy25) / fy25


def promotion_raises(rows: list[dict]) -> dict:
    out = {}
    for to_level, from_level in LEVEL_TO_PREV.items():
        subset = [
            raise_pct(r) for r in rows
            if r.get("FY26 Level") == to_level and r.get("FY25 Level") == from_level
        ]
        subset = [x for x in subset if x is not None and -0.5 < x < 1.0]
        if len(subset) >= 10:
            f_lbl, t_lbl = FROM_LABEL[to_level]
            out[to_level] = {
                "n": len(subset),
                "median": round(float(np.median(subset)), 4),
                "fromLabel": f_lbl,
                "toLabel": t_lbl,
            }
    return out


def non_promotion_raise(rows: list[dict]) -> dict:
    subset = [
        raise_pct(r) for r in rows
        if r.get("FY25 Level") == r.get("FY26 Level")
    ]
    subset = [x for x in subset if x is not None and -0.5 < x < 1.0]
    return {
        "n": len(subset),
        "median": round(float(np.median(subset)), 4),
        "p25": round(float(np.percentile(subset, 25)), 4),
        "p75": round(float(np.percentile(subset, 75)), 4),
    }


def main() -> None:
    raw = load_rows()
    rows = clean(raw)
    print(f"Raw: {len(raw)}  Clean: {len(rows)}", file=sys.stderr)
    assert 1500 <= len(rows) <= 1900, f"Unexpected clean row count: {len(rows)}"
    level = level_stats(rows)
    print(json.dumps({"LEVEL_STATS": level}, indent=2), file=sys.stderr)
    biz = business_stats(rows)
    gps = gps_commercial_stats(rows)
    port = portfolio_stats(rows)
    print(f"Businesses per level: {[len(biz[l]) for l in LEVEL_ORDER]}", file=sys.stderr)
    print(f"GPS/Comm per level:   {[len(gps[l]) for l in LEVEL_ORDER]}", file=sys.stderr)
    print(f"Portfolios per level: {[len(port.get(l, {})) for l in LEVEL_ORDER]}", file=sys.stderr)
    usdc = usdc_stats(rows)
    mba = mba_stats(rows)
    premium = mba_premium(rows)
    years = years_at_level_manager(rows)
    print(f"USDC levels: {list(usdc)}", file=sys.stderr)
    print(f"MBA levels:  {list(mba)}", file=sys.stderr)
    print(f"Premium:     {list(premium)}", file=sys.stderr)
    print(f"Years(mgr):  {years}", file=sys.stderr)
    promos = promotion_raises(rows)
    non_promo = non_promotion_raise(rows)
    print(f"Promos: {promos}", file=sys.stderr)
    print(f"Non-promo: {non_promo}", file=sys.stderr)
    write_js({
        "count": len(rows),
        "level": level,
        "gps": gps,
        "biz": biz,
        "port": port,
        "usdc": usdc,
        "mba": mba,
        "premium": premium,
        "years": years,
        "promos": promos,
        "non_promo": non_promo,
    })
    print(f"Wrote {OUT_PATH} ({len(rows)} respondents)", file=sys.stderr)


if __name__ == "__main__":
    main()
